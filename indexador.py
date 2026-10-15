# -*- coding: utf-8 -*-
"""
Leitura dos Excel da extração do legado e criação da base de consulta (SQLite).

A base é criada uma única vez (ou quando você mandar atualizar). Depois disso o
programa abre em segundos, sem precisar ler os Excel de novo.
"""
from __future__ import annotations

import base64
import json
import os
import zlib
import re
import shutil
import sqlite3
import tempfile
import time
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

VERSAO_SCHEMA = 2
PASTA_INDICE = "_indice"
NOME_DB = "legado.db"


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #
def sem_acento(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    ).lower()


def txt(v) -> str:
    """Converte qualquer valor de célula em texto limpo."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, datetime):
        return v.isoformat()
    s = str(v).replace("_x000D_", "").replace("\r\n", "\n").replace("\r", "\n")
    return s.strip()


def so_digitos(s: str) -> str:
    return re.sub(r"\D", "", s or "")


def chave_num(s: str) -> str:
    """'0171341' -> '171341' (para comparar NFs / remessas sem zeros à esquerda)."""
    return so_digitos(s).lstrip("0")


def norm_caso(v) -> str:
    s = txt(v)
    return s.zfill(8) if s.isdigit() else s


def tira_html(s: str) -> str:
    s = re.sub(r"<br\s*/?>", " · ", s or "", flags=re.I)
    return re.sub(r"<[^>]+>", "", s).strip(" ·")


def nome_semestre(nome_arquivo: str) -> str:
    """'09 - Extração - Segundo Semestre 2026.xlsx' -> '09 - Segundo Semestre 2026'."""
    base = Path(nome_arquivo).stem
    m = re.match(r"^(\d+)\s*-\s*Extra[cç][aã]o\s*-\s*(.+)$", base, re.I)
    return f"{m.group(1)} - {m.group(2).strip()}" if m else base


def pasta_local_dados() -> Path:
    base = os.environ.get("LOCALAPPDATA") or os.path.join(Path.home(), ".local", "share")
    p = Path(base) / "ConsultaLegado"
    p.mkdir(parents=True, exist_ok=True)
    return p


def pasta_indice_compartilhada(root: str | Path) -> Path:
    return Path(root) / PASTA_INDICE


def db_local() -> Path:
    """A base que o programa realmente abre: SEMPRE uma cópia no disco do próprio computador."""
    return pasta_local_dados() / NOME_DB


def db_compartilhada(root: str | Path) -> Path:
    return pasta_indice_compartilhada(root) / NOME_DB


def _arq_origem() -> Path:
    return pasta_local_dados() / "origem.json"


def _arq_versao(root: str | Path) -> Path:
    return pasta_indice_compartilhada(root) / "versao_schema.txt"


def compartilhada_compativel(root: str | Path) -> bool:
    """A base da pasta compartilhada foi criada por esta versão do programa? (não abre o arquivo grande)"""
    try:
        return _arq_versao(root).read_text().strip() == str(VERSAO_SCHEMA)
    except OSError:
        return False


def assinatura(p: Path) -> dict:
    st = p.stat()
    return {"tamanho": st.st_size, "mtime": int(st.st_mtime)}


def copia_local_desatualizada(compartilhada: Path) -> bool:
    """True se a base da pasta compartilhada é diferente da cópia local."""
    if not db_local().exists():
        return True
    if not _arq_origem().exists():
        return False  # base criada neste computador e ainda não publicada
    try:
        gravada = json.loads(_arq_origem().read_text(encoding="utf-8"))
        return gravada != assinatura(compartilhada)
    except (OSError, ValueError):
        return True


def _copiar(origem: Path, destino: Path, progresso, f0: float, f1: float, rotulo: str):
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_suffix(".tmp")
    total = max(origem.stat().st_size, 1)
    feito = 0
    with open(origem, "rb") as fi, open(tmp, "wb") as fo:
        while True:
            bloco = fi.read(8 * 1024 * 1024)
            if not bloco:
                break
            fo.write(bloco)
            feito += len(bloco)
            progresso(f"{rotulo} ({feito / 1e6:,.0f} de {total / 1e6:,.0f} MB)".replace(",", "."),
                      f0 + (f1 - f0) * feito / total)
    os.replace(tmp, destino)


def copiar_para_local(compartilhada: Path, progresso=lambda m, f: None) -> str:
    """Traz a base da pasta compartilhada (OneDrive) para o disco local — uma vez só."""
    sig = assinatura(compartilhada)
    _copiar(compartilhada, db_local(), progresso, 0.0, 1.0, "Copiando a base para o seu computador")
    _arq_origem().write_text(json.dumps(sig), encoding="utf-8")
    return "Base copiada para o seu computador."


def publicar_na_pasta(root: str | Path, progresso=lambda m, f: None, f0=0.0, f1=1.0) -> bool:
    """Grava a base recém-criada na pasta do legado, para os colegas copiarem. Falhar aqui não é grave."""
    try:
        destino = db_compartilhada(root)
        destino.parent.mkdir(parents=True, exist_ok=True)
        _copiar(db_local(), destino, progresso, f0, f1, "Publicando a base na pasta do legado")
        _arq_versao(root).write_text(str(VERSAO_SCHEMA))
        _arq_origem().write_text(json.dumps(assinatura(destino)), encoding="utf-8")
        return True
    except OSError:
        return False


# --------------------------------------------------------------------------- #
# Leitura do Excel (calamine = muito rápido; openpyxl = alternativa)
# --------------------------------------------------------------------------- #
class Planilha:
    def __init__(self, caminho: Path):
        try:
            from python_calamine import CalamineWorkbook

            self._wb = CalamineWorkbook.from_path(str(caminho))
            self._modo = "calamine"
            self.nomes = list(self._wb.sheet_names)
        except ImportError:
            from openpyxl import load_workbook

            self._wb = load_workbook(str(caminho), read_only=True, data_only=True)
            self._modo = "openpyxl"
            self.nomes = list(self._wb.sheetnames)

    def registros(self, nome: str):
        if self._modo == "calamine":
            it = self._wb.get_sheet_by_name(nome).iter_rows()
        else:
            it = self._wb[nome].iter_rows(values_only=True)
        cab = None
        for linha in it:
            if cab is None:
                cab = [txt(c) or None for c in linha]
                continue
            yield {c: v for c, v in zip(cab, linha) if c}

    def fechar(self):
        if self._modo == "openpyxl":
            self._wb.close()

    def achar_aba(self, *candidatas: str) -> str | None:
        norm = {sem_acento(n).strip(): n for n in self.nomes}
        for c in candidatas:
            if c in norm:
                return norm[c]
        for c in candidatas:
            for k, v in norm.items():
                if c in k:
                    return v
        return None


# --------------------------------------------------------------------------- #
# Esquema
# --------------------------------------------------------------------------- #
SCHEMA = """
CREATE TABLE meta (chave TEXT PRIMARY KEY, valor TEXT);
CREATE TABLE semestres (nome TEXT PRIMARY KEY, pasta TEXT);
CREATE TABLE casos (
    numero TEXT PRIMARY KEY, case_id TEXT, semestre TEXT, criado TEXT,
    transportadora TEXT, cnpj TEXT, cnpj_num TEXT,
    status TEXT, motivo TEXT, submotivo TEXT,
    placa TEXT, placa_norm TEXT, motorista TEXT,
    parent_id TEXT, dados BLOB
);
CREATE TABLE comentarios (numero TEXT, criado TEXT, autor TEXT, corpo TEXT, publico INTEGER);
CREATE TABLE nfs (
    numero TEXT, nf TEXT, nf_chave TEXT, serie TEXT, serie_chave TEXT,
    remessa TEXT, remessa_chave TEXT, cliente TEXT, centro TEXT, status_nf TEXT, vendedor TEXT
);
CREATE TABLE anexos (
    numero TEXT, doc_id TEXT, arquivo TEXT, nome TEXT, ext TEXT, tamanho INTEGER,
    criado TEXT, criado_por TEXT, mais_recente INTEGER, versao TEXT,
    compartilhado TEXT, semestre TEXT, achado INTEGER
);
"""

INDICES = """
CREATE INDEX ix_nf ON nfs(nf_chave, serie_chave);
CREATE INDEX ix_remessa ON nfs(remessa_chave);
CREATE INDEX ix_nf_caso ON nfs(numero);
CREATE INDEX ix_cnpj ON casos(cnpj_num);
CREATE INDEX ix_placa ON casos(placa_norm);
CREATE INDEX ix_parent ON casos(parent_id);
CREATE INDEX ix_case_id ON casos(case_id);
CREATE INDEX ix_com ON comentarios(numero);
CREATE INDEX ix_anx ON anexos(numero);
"""


def montar_zdict(amostra: list[dict]) -> bytes:
    """Dicionário de compressão (nomes de campos e valores mais comuns) — reduz a base pela metade."""
    from collections import Counter

    cont: Counter = Counter()
    for r in amostra:
        for k, v in r.items():
            cont[(k, v)] += 1
    chaves = sorted({k for r in amostra for k in r})
    cab = ",".join(f'"{k}":""' for k in chaves)
    minimo = max(3, len(amostra) // 10)
    comuns = ",".join(f'"{k}":"{v}"' for (k, v), n in cont.most_common(400) if n >= minimo and len(v) < 60)
    return (comuns + "," + cab).encode("utf-8")[-32768:]


def comprimir_json(d: dict, zdict: bytes) -> bytes:
    bruto = json.dumps(d, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    o = zlib.compressobj(6, zdict=zdict)
    return o.compress(bruto) + o.flush()


def descomprimir_json(blob: bytes, zdict: bytes) -> dict:
    o = zlib.decompressobj(zdict=zdict)
    return json.loads(o.decompress(blob).decode("utf-8"))


def _amostra_casos(excels: list[Path], n: int = 300) -> list[dict]:
    for x in excels:
        try:
            pl = Planilha(x)
            aba = pl.achar_aba("casos")
            if not aba:
                continue
            res = []
            for r in pl.registros(aba):
                d = {k: txt(v) for k, v in r.items() if k != "_"}
                res.append({k: v for k, v in d.items() if v != ""})
                if len(res) >= n:
                    break
            pl.fechar()
            if res:
                return res
        except Exception:
            continue
    return []


# --------------------------------------------------------------------------- #
# Indexação de um semestre
# --------------------------------------------------------------------------- #
def _achar_pasta_anexos(root: Path, semestre: str) -> Path | None:
    anexos = root / "ANEXOS"
    if not anexos.is_dir():
        return None
    p = anexos / semestre
    if p.is_dir():
        return p
    prefixo = semestre.split(" - ")[0].strip()
    for e in anexos.iterdir():
        if e.is_dir() and e.name.split(" - ")[0].strip() == prefixo:
            return e
    return None


def _listar_arquivos(pasta: Path | None):
    """Devolve (pasta_real, {doc_id: [nomes]}, conjunto de nomes)."""
    por_doc: dict[str, list[str]] = defaultdict(list)
    nomes: set[str] = set()
    if pasta is None:
        return None, por_doc, nomes
    real = pasta / "arquivos" if (pasta / "arquivos").is_dir() else pasta
    with os.scandir(real) as it:
        for e in it:
            if not e.is_file():
                continue
            nomes.add(e.name)
            m = re.match(r"^[^_]*_(069[0-9A-Za-z]{12,15})_", e.name)
            if m:
                por_doc[m.group(1)].append(e.name)
    return real, por_doc, nomes


def _indexar_arquivo(conn, xlsx: Path, root: Path, estado: dict, cb) -> dict:
    sem = nome_semestre(xlsx.name)
    cb(f"[{sem}] abrindo planilha…", 0.0)
    pl = Planilha(xlsx)
    aba_casos = pl.achar_aba("casos")
    aba_com = pl.achar_aba("comentarios do caso", "comentarios")
    aba_nf = pl.achar_aba("nf-caso", "nf caso")
    aba_link = pl.achar_aba("contentdocumentlink")
    aba_ver = pl.achar_aba("contentversion")
    stats = {"semestre": sem, "casos": 0, "comentarios": 0, "nfs": 0, "anexos": 0,
             "anexos_sem_arquivo": 0, "arquivos_orfaos": 0, "pasta_anexos": None,
             "avisos": []}
    if not aba_casos:
        stats["avisos"].append("aba CASOS não encontrada — arquivo ignorado")
        pl.fechar()
        return stats

    mapa_id: dict[str, str] = {}

    # ---- CASOS ----------------------------------------------------------- #
    cb(f"[{sem}] lendo casos…", 0.05)
    principais = []
    for r in pl.registros(aba_casos):
        cn = norm_caso(r.get("CaseNumber"))
        if not cn:
            continue
        d = {k: txt(v) for k, v in r.items() if k != "_"}
        d = {k: v for k, v in d.items() if v != ""}
        cid = d.get("Id", "")
        mapa_id[cid] = cn
        lm = d.get("LastModifiedDate", "")
        anterior = estado["ultima_mod"].get(cn)
        if anterior is not None and anterior >= lm:
            continue
        estado["ultima_mod"][cn] = lm
        placa = d.get("PlacaVeiculo__c", "")
        motivo = d.get("MotivoView__c") or d.get("Motivo__c") or d.get("Classificacao_Interna_Motivo__c", "")
        submotivo = (d.get("SubmotivoView__c") or d.get("Submotivo__c")
                     or d.get("Classificacao_Interna_Submotivo__c", ""))
        dados = comprimir_json(d, estado["zdict"])
        conn.execute(
            "INSERT OR REPLACE INTO casos VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                cn, cid, sem, d.get("CreatedDate", ""),
                d.get("Account.Name", ""), d.get("Account.TaxNumber__c", ""),
                so_digitos(d.get("Account.TaxNumber__c", "")),
                d.get("Status", ""), motivo, submotivo,
                placa, re.sub(r"[^0-9A-Za-z]", "", placa).upper(),
                d.get("NomeMotorista__c", ""), d.get("ParentId", ""), dados,
            ),
        )
        # texto para a busca livre (fica numa tabela temporária, fora do arquivo final)
        conn.execute("INSERT OR REPLACE INTO temp.t_desc VALUES (?,?)",
                     (cn, f"{d.get('Subject', '')} {d.get('Description', '')}"))
        stats["casos"] += 1
        if d.get("NumeroNF__c"):
            principais.append((cn, d.get("NumeroNF__c", ""), d.get("SerieNF__c", "")))

    # ---- COMENTÁRIOS ----------------------------------------------------- #
    if aba_com:
        cb(f"[{sem}] lendo comentários…", 0.25)
        for r in pl.registros(aba_com):
            cn = norm_caso(r.get("Parent.CaseNumber")) or mapa_id.get(txt(r.get("ParentId")), "")
            cid = txt(r.get("Id"))
            if not cn or not cid or cid in estado["com"]:
                continue
            estado["com"].add(cid)
            conn.execute(
                "INSERT INTO comentarios VALUES (?,?,?,?,?)",
                (cn, txt(r.get("CreatedDate")), txt(r.get("CreatedBy.Name")), txt(r.get("CommentBody")),
                 0 if txt(r.get("IsPublished")).lower() == "false" else 1),
            )
            stats["comentarios"] += 1
    else:
        stats["avisos"].append("aba de comentários não encontrada")

    # ---- NFs ------------------------------------------------------------- #
    cb(f"[{sem}] lendo NFs…", 0.45)
    ja_tem = set()
    if aba_nf:
        for r in pl.registros(aba_nf):
            cn = norm_caso(r.get("NumeroCaso__c")) or mapa_id.get(txt(r.get("Caso__c")), "")
            rid = txt(r.get("Id"))
            nf = txt(r.get("NumeroNota__c"))
            if not nf:
                nf = tira_html(txt(r.get("NumNotaFiscal__c"))).split("-")[0]
            if not cn or not rid or rid in estado["nf"]:
                continue
            estado["nf"].add(rid)
            serie = txt(r.get("SerieNota__c"))
            remessa = txt(r.get("Remessa__c"))
            conn.execute(
                "INSERT INTO nfs VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                (cn, nf, chave_num(nf), serie, chave_num(serie), remessa, chave_num(remessa),
                 txt(r.get("Cliente__c")), txt(r.get("Centro_DT__c")), txt(r.get("Status__c")),
                 tira_html(txt(r.get("Vendedor__c")))),
            )
            stats["nfs"] += 1
            ja_tem.add((cn, chave_num(nf)))
    else:
        stats["avisos"].append("aba NF-CASO não encontrada")
    # NF informada no próprio caso e que não está na aba NF-CASO
    for cn, nf, serie in principais:
        if (cn, chave_num(nf)) in ja_tem or (cn, chave_num(nf)) in estado["nf_caso"]:
            continue
        estado["nf_caso"].add((cn, chave_num(nf)))
        conn.execute(
            "INSERT INTO nfs VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (cn, nf, chave_num(nf), serie, chave_num(serie), "", "", "", "", "", ""),
        )
        stats["nfs"] += 1

    # ---- ANEXOS ---------------------------------------------------------- #
    if aba_ver:
        cb(f"[{sem}] lendo anexos…", 0.6)
        docs_casos: dict[str, list[str]] = defaultdict(list)
        if aba_link:
            for r in pl.registros(aba_link):
                doc = txt(r.get("ContentDocumentId"))
                cn = mapa_id.get(txt(r.get("LinkedEntityId")))
                if doc and cn and cn not in docs_casos[doc]:
                    docs_casos[doc].append(cn)
        else:
            stats["avisos"].append("aba CONTENTDOCUMENTLINK não encontrada")

        pasta = _achar_pasta_anexos(root, sem)
        pasta_real, por_doc, nomes = _listar_arquivos(pasta)
        stats["pasta_anexos"] = str(pasta_real) if pasta_real else None
        if pasta_real is None:
            stats["avisos"].append("pasta de anexos do semestre não encontrada em ANEXOS/")
        rel_base = (
            pasta_real.relative_to(root).as_posix()
            if pasta_real else f"ANEXOS/{sem}/arquivos"
        )
        conn.execute("INSERT OR REPLACE INTO semestres VALUES (?,?)", (sem, rel_base))
        docs_vistos = set()
        for r in pl.registros(aba_ver):
            doc = txt(r.get("ContentDocumentId"))
            if not doc:
                continue
            docs_vistos.add(doc)
            nome = txt(r.get("PathOnClient"))
            caso_nome = norm_caso(r.get("CASENUMBER")) or (docs_casos.get(doc) or [""])[0]
            esperado = f"{caso_nome}_{doc}_{nome}"
            achado = 1
            if esperado in nomes:
                arquivo = esperado
            elif por_doc.get(doc):
                cands = por_doc[doc]
                arquivo = next((c for c in cands if c.endswith("_" + nome)), cands[0])
            else:
                arquivo, achado = esperado, 0
            casos_do_doc = docs_casos.get(doc) or ([caso_nome] if caso_nome else [])
            tamanho = txt(r.get("ContentSize"))
            vid = txt(r.get("Id"))
            for cn in casos_do_doc:
                if (cn, vid) in estado["anx"]:
                    continue
                estado["anx"].add((cn, vid))
                outros = ", ".join(c for c in casos_do_doc if c != cn)
                conn.execute(
                    "INSERT INTO anexos VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (cn, doc, arquivo, nome,
                     (txt(r.get("FileExtension")) or Path(nome).suffix.lstrip(".")).lower(),
                     int(tamanho) if tamanho.isdigit() else 0,
                     txt(r.get("CreatedDate")), txt(r.get("CreatedBy.Name")),
                     0 if txt(r.get("IsLatest")).lower() == "false" else 1,
                     txt(r.get("VersionNumber")), outros, sem, achado),
                )
                stats["anexos"] += 1
                if not achado:
                    stats["anexos_sem_arquivo"] += 1
        stats["arquivos_orfaos"] = sum(1 for d in por_doc if d not in docs_vistos)
    else:
        stats["avisos"].append("aba CONTENTVERSION não encontrada")

    pl.fechar()
    return stats


# --------------------------------------------------------------------------- #
# Construção da base completa
# --------------------------------------------------------------------------- #
def listar_excels(root: Path) -> list[Path]:
    return sorted(
        p for p in root.glob("*.xlsx")
        if not p.name.startswith("~$") and re.match(r"^\d+", p.name)
    )


def construir_base(root: str | Path, progresso=lambda msg, frac: None) -> tuple[str, bool]:
    """Lê todos os Excel da pasta e cria a base NO DISCO LOCAL; depois tenta publicá-la na pasta do
    legado para os colegas. Devolve (relatório, publicou_na_pasta)."""
    root = Path(root)
    excels = listar_excels(root)
    if not excels:
        raise ValueError(
            "Nenhum arquivo Excel numerado (ex.: '01 - Extração - 2022.xlsx') foi encontrado "
            f"em:\n{root}"
        )
    t0 = time.time()
    tmpdir = Path(tempfile.mkdtemp(prefix="consulta_legado_"))
    tmp_db = tmpdir / NOME_DB
    conn = sqlite3.connect(tmp_db)
    conn.execute("PRAGMA journal_mode=OFF")
    conn.execute("PRAGMA synchronous=OFF")
    conn.executescript(SCHEMA)
    conn.execute("CREATE TEMP TABLE t_desc (numero TEXT PRIMARY KEY, texto TEXT)")
    progresso("Preparando a compressão dos dados…", 0.0)
    zdict = montar_zdict(_amostra_casos(excels))
    estado = {"ultima_mod": {}, "com": set(), "nf": set(), "nf_caso": set(), "anx": set(), "zdict": zdict}
    todas = []
    n = len(excels)
    try:
        for i, x in enumerate(excels):
            st = _indexar_arquivo(
                conn, x, root, estado,
                lambda m, f, i=i: progresso(m, (i + f) / n * 0.85),
            )
            todas.append(st)
            conn.commit()
        estado.clear()
        conn.execute("INSERT INTO meta VALUES ('zdict', ?)", (base64.b64encode(zdict).decode("ascii"),))

        progresso("Criando índices de busca…", 0.86)
        conn.executescript(INDICES)
        conn.execute(
            "CREATE VIRTUAL TABLE busca USING fts5(texto, content='', detail=column, "
            "tokenize='unicode61 remove_diacritics 2')"
        )
        conn.execute(
            """INSERT INTO busca(rowid, texto)
               SELECT c.rowid,
                 c.numero||' '||coalesce(d.texto,'')||' '||
                 coalesce(c.transportadora,'')||' '||coalesce(c.placa,'')||' '||coalesce(c.motorista,'')||' '||
                 coalesce(c.motivo,'')||' '||coalesce(c.submotivo,'')||' '||
                 coalesce((SELECT group_concat(corpo,' ') FROM comentarios WHERE numero=c.numero),'')||' '||
                 coalesce((SELECT group_concat(nome,' ') FROM anexos WHERE numero=c.numero),'')
               FROM casos c LEFT JOIN temp.t_desc d ON d.numero=c.numero"""
        )
        conn.execute("DROP TABLE temp.t_desc")
        progresso("Finalizando…", 0.9)

        tot = {k: sum(s[k] for s in todas)
               for k in ("casos", "comentarios", "nfs", "anexos", "anexos_sem_arquivo", "arquivos_orfaos")}
        linhas = [
            f"Base criada em {datetime.now():%d/%m/%Y %H:%M} ({time.time()-t0:.0f}s)",
            f"Pasta do legado: {root}",
            "",
            f"{'Semestre':<34}{'Casos':>8}{'Coment.':>9}{'NFs':>8}{'Anexos':>8}{'Sem arq.':>9}{'Órfãos':>8}",
        ]
        for s in todas:
            linhas.append(
                f"{s['semestre']:<34}{s['casos']:>8}{s['comentarios']:>9}{s['nfs']:>8}"
                f"{s['anexos']:>8}{s['anexos_sem_arquivo']:>9}{s['arquivos_orfaos']:>8}"
            )
        linhas.append(
            f"{'TOTAL':<34}{tot['casos']:>8}{tot['comentarios']:>9}{tot['nfs']:>8}"
            f"{tot['anexos']:>8}{tot['anexos_sem_arquivo']:>9}{tot['arquivos_orfaos']:>8}"
        )
        linhas += ["", "'Sem arq.' = anexo cadastrado na planilha mas arquivo não encontrado na pasta.",
                   "'Órfãos'  = arquivo na pasta sem registro na planilha."]
        for s in todas:
            for a in s["avisos"]:
                linhas.append(f"AVISO [{s['semestre']}]: {a}")
        relatorio = "\n".join(linhas)

        conn.executemany(
            "INSERT INTO meta VALUES (?,?)",
            [("schema", str(VERSAO_SCHEMA)), ("criado_em", datetime.now().isoformat(timespec="seconds")),
             ("pasta_legado", str(root)), ("relatorio", relatorio),
             ("total_casos", str(tot["casos"]))],
        )
        conn.commit()
        conn.close()

        # 1) grava a base local (é a que o programa usa)
        _copiar(tmp_db, db_local(), progresso, 0.9, 0.93, "Gravando a base no seu computador")
        if _arq_origem().exists():
            _arq_origem().unlink()
        # 2) publica na pasta do legado para os colegas (opcional)
        publicou = publicar_na_pasta(root, progresso, 0.93, 1.0)
        tamanho = db_local().stat().st_size / 1e6
        relatorio += (f"\n\nTamanho da base: {tamanho:,.0f} MB".replace(",", ".")
                      + ("\nBase publicada em: " + str(db_compartilhada(root)) if publicou else
                         "\nA pasta do legado não aceitou gravação: a base ficou só neste computador."))
        try:
            (db_compartilhada(root).parent / "relatorio_carga.txt").write_text(relatorio, encoding="utf-8")
        except OSError:
            pass
        progresso("Concluído", 1.0)
        return relatorio, publicou
    finally:
        try:
            conn.close()
        except Exception:
            pass
        shutil.rmtree(tmpdir, ignore_errors=True)
